
from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import load_workbook
from io import BytesIO
import re

app = Flask(__name__)
CORS(app)


def normalizar_servico(valor):
    """Padroniza o número do serviço para permitir a comparação."""
    if valor is None:
        return None

    texto = str(valor).strip()

    if not texto:
        return None

    # Corrige números que o Excel entrega como 40205.0
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]

    return texto


def adicionar_servico(base, valor, linha, coluna):
    """Registra um serviço e a localização em que ele aparece."""
    servico = normalizar_servico(valor)

    if not servico:
        return

    if servico not in base:
        base[servico] = []

    base[servico].append({
        "linha": linha,
        "coluna": coluna
    })


def separar_servicos(valor):
    """Separa serviços que estejam juntos, como 40205/40206."""
    if valor is None:
        return []

    texto = str(valor).strip()

    if not texto:
        return []

    partes = re.split(r"\s*/\s*", texto)

    return [
        normalizar_servico(parte)
        for parte in partes
        if normalizar_servico(parte)
    ]


@app.route("/api/status", methods=["GET"])
def status():
    return jsonify({
        "mensagem": "Robô Nietzsche está funcionando!"
    })


@app.route("/api/comparar", methods=["POST"])
def comparar():

    programacao = request.files.get("programacao")
    otimizacao = request.files.get("otimizacao")

    if not programacao or not otimizacao:
        return jsonify({
            "erro": "Envie as duas planilhas."
        }), 400

    try:
        planilha_programacao = load_workbook(
            BytesIO(programacao.read()),
            data_only=True
        )

        planilha_otimizacao = load_workbook(
            BytesIO(otimizacao.read()),
            data_only=True
        )

        # Conferir se as abas necessárias existem
        if "Demanda" not in planilha_programacao.sheetnames:
            raise ValueError(
                "A aba 'Demanda' não foi encontrada na Alta Demanda."
            )

        if "Cenario" not in planilha_otimizacao.sheetnames:
            raise ValueError(
                "A aba 'Cenario' não foi encontrada na Otimização."
            )

        aba_demanda = planilha_programacao["Demanda"]
        aba_cenario = planilha_otimizacao["Cenario"]

        # Dicionários com os serviços e suas localizações
        servicos_demanda = {}
        servicos_otimizacao = {}

        # ---------------------------------
        # LER ALTA DEMANDA
        # Q = TIPO
        # D = SERVIÇO 1
        # E = SERVIÇO 2
        # ---------------------------------

        for linha in range(3, aba_demanda.max_row + 1):

            tipo = aba_demanda.cell(
                row=linha,
                column=17
            ).value

            if str(tipo).strip().upper() != "HO":
                continue

            servico_1 = aba_demanda.cell(
                row=linha,
                column=4
            ).value

            servico_2 = aba_demanda.cell(
                row=linha,
                column=5
            ).value

            adicionar_servico(
                servicos_demanda,
                servico_1,
                linha,
                "D"
            )

            adicionar_servico(
                servicos_demanda,
                servico_2,
                linha,
                "E"
            )

        # ---------------------------------
        # LER OTIMIZAÇÃO
        # B = SERVIÇO
        # F = TIPO
        # ---------------------------------

        for linha in range(3, aba_cenario.max_row + 1):

            tipo = aba_cenario.cell(
                row=linha,
                column=6
            ).value

            if str(tipo).strip().upper() not in ["HO", "EHO"]:
                continue

            valor_servico = aba_cenario.cell(
                row=linha,
                column=2
            ).value

            for servico in separar_servicos(valor_servico):

                adicionar_servico(
                    servicos_otimizacao,
                    servico,
                    linha,
                    "B"
                )

        # ---------------------------------
        # COMPARAR AS DUAS PLANILHAS
        # ---------------------------------

        conjunto_demanda = set(servicos_demanda.keys())
        conjunto_otimizacao = set(servicos_otimizacao.keys())

        correspondentes = conjunto_demanda & conjunto_otimizacao

        ausentes_otimizacao = (
            conjunto_demanda - conjunto_otimizacao
        )

        ausentes_demanda = (
            conjunto_otimizacao - conjunto_demanda
        )

        # Organizar os resultados para o site
        resultado_correspondentes = []

        for servico in sorted(correspondentes):
            resultado_correspondentes.append({
                "servico": servico,
                "linhas_demanda": servicos_demanda[servico],
                "linhas_otimizacao": servicos_otimizacao[servico]
            })

        resultado_ausentes_otimizacao = []

        for servico in sorted(ausentes_otimizacao):
            resultado_ausentes_otimizacao.append({
                "servico": servico,
                "linhas": servicos_demanda[servico]
            })

        resultado_ausentes_demanda = []

        for servico in sorted(ausentes_demanda):
            resultado_ausentes_demanda.append({
                "servico": servico,
                "linhas": servicos_otimizacao[servico]
            })

        return jsonify({
            "mensagem": "Comparação concluída com sucesso!",

            "resumo": {
                "total_demanda": len(conjunto_demanda),
                "total_otimizacao": len(conjunto_otimizacao),
                "correspondentes": len(correspondentes),
                "ausentes_otimizacao": len(ausentes_otimizacao),
                "ausentes_demanda": len(ausentes_demanda)
            },

            "correspondentes": resultado_correspondentes,

            "ausentes_otimizacao": resultado_ausentes_otimizacao,

            "ausentes_demanda": resultado_ausentes_demanda
        })

    except Exception as erro:
        return jsonify({
            "erro": f"Não foi possível comparar as planilhas: {erro}"
        }), 400


if __name__ == "__main__":
    app.run(debug=True, port=5000)
